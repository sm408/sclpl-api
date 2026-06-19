from __future__ import annotations

import csv
import json
import uuid
from pathlib import Path

from app.core.contracts.export_pipeline import ExportPipeline, ExportResult
from app.core.models.export import ExportFormat


class DefaultExportPipeline(ExportPipeline):
    async def export_json(self, data: list[dict], output_path: str) -> ExportResult:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, default=str)
        return ExportResult(
            path=str(path),
            format=ExportFormat.JSON,
            record_count=len(data),
            success=True,
        )

    async def export_csv(self, data: list[dict], output_path: str) -> ExportResult:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if not data:
            return ExportResult(
                path=str(path),
                format=ExportFormat.CSV,
                record_count=0,
                success=True,
            )
        fieldnames = list(data[0].keys())
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)
        return ExportResult(
            path=str(path),
            format=ExportFormat.CSV,
            record_count=len(data),
            success=True,
        )

    async def export_excel(self, data: list[dict], output_path: str) -> ExportResult:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            return ExportResult(
                path=str(path),
                format=ExportFormat.EXCEL,
                record_count=0,
                success=False,
                error="openpyxl not installed. Install with: pip install sclplapi[excel]",
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Export"

        if not data:
            wb.save(str(path))
            return ExportResult(
                path=str(path),
                format=ExportFormat.EXCEL,
                record_count=0,
                success=True,
            )

        headers = list(data[0].keys())
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3A7BD5", end_color="3A7BD5", fill_type="solid")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, "")
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, default=str)
                ws.cell(row=row_idx, column=col_idx, value=value)

        for col_idx, header in enumerate(headers, 1):
            max_len = max(
                len(str(header)),
                max((len(str(row.get(header, ""))) for row in data), default=0),
            )
            ws.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 50)

        wb.save(str(path))
        return ExportResult(
            path=str(path),
            format=ExportFormat.EXCEL,
            record_count=len(data),
            success=True,
        )
