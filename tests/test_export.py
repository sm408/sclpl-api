from app.services.export_service import DefaultExportPipeline
import json
import csv
import pytest
from pathlib import Path


@pytest.mark.asyncio
async def test_export_json(tmp_path):
    pipeline = DefaultExportPipeline()
    data = [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]
    output = str(tmp_path / "output.json")
    result = await pipeline.export_json(data, output)
    assert result.success
    assert result.record_count == 2
    with open(output) as f:
        loaded = json.load(f)
    assert len(loaded) == 2
    assert loaded[0]["name"] == "Alice"


@pytest.mark.asyncio
async def test_export_csv(tmp_path):
    pipeline = DefaultExportPipeline()
    data = [{"name": "Alice", "age": "30"}, {"name": "Bob", "age": "25"}]
    output = str(tmp_path / "output.csv")
    result = await pipeline.export_csv(data, output)
    assert result.success
    assert result.record_count == 2
    with open(output) as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    assert len(rows) == 2
    assert rows[0]["name"] == "Alice"


@pytest.mark.asyncio
async def test_export_csv_empty(tmp_path):
    pipeline = DefaultExportPipeline()
    output = str(tmp_path / "empty.csv")
    result = await pipeline.export_csv([], output)
    assert result.success
    assert result.record_count == 0


@pytest.mark.asyncio
async def test_export_excel(tmp_path):
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    pipeline = DefaultExportPipeline()
    data = [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]
    output = str(tmp_path / "output.xlsx")
    result = await pipeline.export_excel(data, output)
    assert result.success
    assert result.record_count == 2
    assert Path(output).exists()

    wb = openpyxl.load_workbook(output)
    ws = wb.active
    assert ws.cell(1, 1).value == "name"
    assert ws.cell(2, 1).value == "Alice"
    assert ws.cell(3, 1).value == "Bob"


@pytest.mark.asyncio
async def test_export_excel_empty(tmp_path):
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    pipeline = DefaultExportPipeline()
    output = str(tmp_path / "empty.xlsx")
    result = await pipeline.export_excel([], output)
    assert result.success
    assert result.record_count == 0
